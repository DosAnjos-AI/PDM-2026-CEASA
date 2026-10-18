#!/usr/bin/env python3
"""
Coletor das cotacoes diarias da CEASA-MS.

O site publica, por ano (pagina boletim-AAAA/), uma planilha Excel por mes.
Cada planilha tem uma aba por dia do mes ("1" a "31") com a tabela de precos
dividida em grupos (AVES, OVOS / CARNES / FRUTAS / HORTALICAS / ...).

Duas etapas, ambas idempotentes:
  1. coleta: baixa as planilhas para <bruto>/AAAA/MM.xls. A fonte preenche a
     planilha do mes aos poucos e a sobe de novo com o mesmo nome, entao as
     dos ultimos --meses-recentes meses sao sempre baixadas de novo e gravadas
     se o conteudo mudou: o tamanho nao serve de sinal, porque o .xls cresce em
     blocos fixos e ja reserva as 31 abas. As mais antigas so sao baixadas de
     novo se o tamanho remoto (HEAD) difere do local.
  2. extracao: le todas as abas e gera o CSV no esquema comum. Com
     --carregados, so as planilhas que ainda nao estao na Bronze e as dos
     ultimos --meses-recentes (a do mes corrente sempre entra).

Particularidade da fonte, verificada em ago/2026: todas as 31 abas existem e
estao preenchidas, inclusive sabados, domingos e dias sem pregao, que repetem
a aba anterior. O extrator reproduz a fonte como esta (a data e a da celula
"Data:" de cada aba); descartar os dias repetidos e decisao da Silver.
Abas cuja data nao bate com o nome da aba e o mes do arquivo, ou que estao no
futuro, sao listadas e ficam fora do CSV.

Dependencias: requests, beautifulsoup4, xlrd (.xls) e openpyxl (.xlsx);
msoffcrypto-tool opcional, so para planilhas com protecao contra gravacao.

Uso:
    python3 ingestao/scraping/ceasa_ms.py                   # 2021 ate hoje
    python3 ingestao/scraping/ceasa_ms.py --start-year 2026
    python3 ingestao/scraping/ceasa_ms.py --so-extrair
"""

from __future__ import annotations

import argparse
import io
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comum import (DIR_DADOS, Cliente, argumentos_incrementais, grava_atomico, grava_se_conferido,  # noqa: E402
                   grava_vazio, http_404, limpa, linha, mes_recente, numero_br, seleciona_pendentes)

# ------------------------------------------------------------ configuracao

URL_ANO = "https://www.ceasa.ms.gov.br/boletim-{ano}/"
FONTE = "CEASA-MS"
MERCADO = "CAMPO GRANDE"
ANO_INICIAL_PADRAO = 2021  # boletim-2020/ e anteriores: 404
BRUTO_PADRAO = DIR_DADOS / "cotacoes_ms"
SAIDA_PADRAO = DIR_DADOS / "cotacoes_ms.csv"
EXTENSOES = (".xls", ".xlsx")
FIM_DA_TABELA = "SIGLAS"  # a partir daqui a aba so tem legenda
# senha publica que o Excel usa na "protecao contra gravacao" (ex.: 2021/04.xls)
SENHA_PADRAO_EXCEL = "VelvetSweatshop"
# preco gravado como texto com ponto decimal (ex.: '29.90')
RE_PONTO_DECIMAL = re.compile(r"^\d+\.\d{1,2}$")

MESES = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
    "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}


def sem_acento(texto: str) -> str:
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c)).lower()


# ---------------------------------------------------------------- coleta

def planilhas_do_ano(cli: Cliente, ano: int) -> dict[int, str] | None:
    """{mes: url} da pagina do ano, ou None se a pagina nao existe."""
    url = URL_ANO.format(ano=ano)
    try:
        r = cli.get(url)
    except requests.HTTPError as e:
        if e.response is not None and e.response.status_code == 404:
            return None
        raise
    achados: dict[int, str] = {}
    for a in BeautifulSoup(r.text, "html.parser").select("a[href]"):
        href = urljoin(url, a["href"])
        if not href.lower().endswith(EXTENSOES):
            continue
        mes = MESES.get(sem_acento(a.get_text(" ", strip=True)))
        if mes is None:
            raise SystemExit(f"{url}: link de planilha com texto que nao e mes: '{a.get_text(strip=True)}' -> {href}")
        if mes in achados and achados[mes] != href:
            raise SystemExit(f"{url}: dois links para o mes {mes}: {achados[mes]} e {href}")
        achados[mes] = href
    return achados


def coletar(cli: Cliente, bruto: Path, inicio: int, fim: int, meses_recentes: int = 2) -> int:
    contagem = {"baixado": 0, "atualizado": 0, "existente": 0, "indisponivel": 0, "erro": 0}
    for ano in range(inicio, fim + 1):
        meses = planilhas_do_ano(cli, ano)
        if meses is None:
            print(f"{ano}: pagina inexistente ({URL_ANO.format(ano=ano)})")
            continue
        print(f"{ano}: {len(meses)} planilha(s)")
        for mes, url in sorted(meses.items()):
            ext = Path(url).suffix.lower()
            destino = bruto / f"{ano:04d}" / f"{mes:02d}{ext}"
            recente = mes_recente(ano, mes, meses_recentes)
            try:
                if destino.exists():
                    if not recente:
                        remoto = cli.head(url).headers.get("Content-Length")
                        if remoto is not None and int(remoto) == destino.stat().st_size:
                            contagem["existente"] += 1
                            continue
                    estado = "atualizado"
                else:
                    estado = "baixado"
                conteudo = cli.get(url).content
            except requests.RequestException as e:
                if http_404(e):
                    # o proprio site publica o link quebrado (ex.: nov/2023, mar/2025):
                    # nao ha o que coletar, e repetir o erro a cada execucao esconderia falha nova
                    contagem["indisponivel"] += 1
                    print(f"  aviso: {ano}-{mes:02d} indisponivel na fonte (404): {url}")
                    continue
                contagem["erro"] += 1
                print(f"  ! {ano}-{mes:02d} {url}: {e}", file=sys.stderr)
                continue
            # .xls e um arquivo OLE2; .xlsx e um zip
            assinatura = b"\xd0\xcf\x11\xe0" if ext == ".xls" else b"PK"
            if not conteudo.startswith(assinatura):
                contagem["erro"] += 1
                print(f"  ! {ano}-{mes:02d}: conteudo nao e {ext} ({url})", file=sys.stderr)
                continue
            # planilha recente baixada de novo e igual a local: nada a gravar
            if estado == "atualizado" and destino.read_bytes() == conteudo:
                contagem["existente"] += 1
                continue
            grava_atomico(destino, conteudo)
            contagem[estado] += 1
            print(f"  {ano}-{mes:02d}: {estado} ({len(conteudo)} bytes)")
    print("Coleta: " + ", ".join(f"{k}={v}" for k, v in contagem.items()))
    return contagem["erro"]


# -------------------------------------------------------------- extracao

class PlanilhaProtegida(Exception):
    pass


def decifra(caminho: Path) -> bytes:
    """Abre planilha com protecao contra gravacao (senha padrao do Excel)."""
    try:
        import msoffcrypto
    except ImportError:
        raise PlanilhaProtegida("protegida; instale msoffcrypto-tool para le-la") from None
    with caminho.open("rb") as fh:
        arq = msoffcrypto.OfficeFile(fh)
        arq.load_key(password=SENHA_PADRAO_EXCEL)
        saida = io.BytesIO()
        arq.decrypt(saida)
    return saida.getvalue()


def abas(caminho: Path):
    """(nome da aba, linhas) com as datas ja como `date`, para .xls e .xlsx."""
    if caminho.suffix.lower() == ".xls":
        import xlrd
        try:
            wb = xlrd.open_workbook(str(caminho))
        except xlrd.XLRDError as e:
            if "encrypted" not in str(e):
                raise
            wb = xlrd.open_workbook(file_contents=decifra(caminho))
        for sh in wb.sheets():
            linhas = []
            for r in range(sh.nrows):
                valores = []
                for c in range(sh.ncols):
                    cel = sh.cell(r, c)
                    if cel.ctype == xlrd.XL_CELL_DATE:
                        valores.append(xlrd.xldate_as_datetime(cel.value, wb.datemode).date())
                    else:
                        valores.append(cel.value)
                linhas.append(valores)
            yield sh.name, linhas
    else:
        import openpyxl
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        for ws in wb.worksheets:
            yield ws.title, [[v.date() if isinstance(v, datetime) else v for v in row]
                             for row in ws.iter_rows(values_only=True)]


def preco(valor) -> float | None:
    """Celula numerica, ou texto: '29,90' ou '29.90' (ponto so com 1-2 casas, sem ambiguidade com milhar)."""
    if isinstance(valor, str) and RE_PONTO_DECIMAL.match(valor.strip()):
        return float(valor.strip())
    return numero_br(valor)


def data_da_aba(linhas: list[list], datemode_xls: bool) -> date | None:
    for valores in linhas:
        if valores and limpa(valores[0]).rstrip(":").lower() == "data":
            for v in valores[1:]:
                if isinstance(v, date):
                    return v
                if isinstance(v, float) and v > 0 and datemode_xls:
                    # celula numerica sem formato de data: serial do Excel
                    import xlrd
                    return xlrd.xldate_as_datetime(v, 0).date()
            return None
    return None


def extrai_planilha(caminho: Path, bruto: Path, descartes: list[str], falhas: list[str]) -> list[dict]:
    ano, mes = int(caminho.parent.name), int(caminho.stem)
    rel = caminho.relative_to(bruto).as_posix()
    hoje = date.today()
    saida: list[dict] = []
    try:
        todas = list(abas(caminho))
    except PlanilhaProtegida as e:
        descartes.append(f"{rel}: {e}")
        return []
    except Exception as e:  # xlrd, openpyxl e msoffcrypto tem hierarquias de erro proprias
        falhas.append(f"{rel}: planilha ilegivel ({e.__class__.__name__}: {e})")
        return []
    for nome, linhas in todas:
        if not nome.strip().isdigit():
            continue  # aba "Menu"
        dia = data_da_aba(linhas, caminho.suffix.lower() == ".xls")
        if dia is None:
            descartes.append(f"{rel} aba {nome}: sem data")
            continue
        if (dia.year, dia.month, dia.day) != (ano, mes, int(nome)):
            descartes.append(f"{rel} aba {nome}: data interna {dia} nao bate com {ano}-{mes:02d}-{int(nome):02d}")
            continue
        if dia > hoje:
            descartes.append(f"{rel} aba {nome}: data futura {dia}")
            continue

        grupo, ordem, depois_da_data = "", 0, False
        for valores in linhas:
            # a tabela ocupa as 6 primeiras colunas; ha sujeira solta mais a direita em algumas abas
            valores = list(valores[:6])
            texto = [limpa(v) for v in valores]
            if not depois_da_data:
                depois_da_data = texto and texto[0].rstrip(":").lower() == "data"
                continue
            if not any(texto):
                continue
            if texto[0].upper() == FIM_DA_TABELA:
                break
            if texto[0] == "Produto":
                continue  # cabecalho repetido a cada grupo
            if texto[0] and not any(texto[2:]):
                # linha so com o nome do grupo (as vezes com sujeira na coluna de embalagem)
                grupo = texto[0]
                continue
            if len(valores) < 6 or not texto[0]:
                raise SystemExit(f"{rel} aba {nome}: linha fora do formato: {valores}")
            saida.append(linha(
                dia, FONTE, MERCADO, grupo, texto[0], f"{rel}#{nome}", ordem,
                embalagem=texto[1], preco_minimo=preco(valores[2]),
                preco_comum=preco(valores[3]), preco_maximo=preco(valores[4]),
                situacao=texto[5],
            ))
            ordem += 1
    return saida


def extrair(bruto: Path, saida: Path, falhas_esperadas: int,
            carregados: Path | None = None, meses_recentes: int = 2) -> int:
    arquivos = sorted(p for p in bruto.rglob("*") if p.suffix.lower() in EXTENSOES)
    # o mesmo mes em .xls e .xlsx (site trocou o formato) duplicaria as linhas do mes
    por_mes: dict[tuple[str, str], list[str]] = {}
    for a in arquivos:
        por_mes.setdefault((a.parent.name, a.stem), []).append(a.relative_to(bruto).as_posix())
    repetidos = [str(v) for v in por_mes.values() if len(v) > 1]
    if repetidos:
        raise SystemExit("Mes com mais de uma planilha no bruto (apague a antiga): " + "; ".join(repetidos))
    arquivos = seleciona_pendentes(arquivos, bruto, carregados, meses_recentes)
    if not arquivos:
        return grava_vazio(saida, FONTE)

    descartes: list[str] = []
    falhas: list[str] = []
    linhas = [l for a in arquivos for l in extrai_planilha(a, bruto, descartes, falhas)]
    print(f"\nExtracao: {len(arquivos)} planilha(s)")
    for titulo, lista in (("Abas fora do CSV", descartes), ("Planilhas ilegiveis", falhas)):
        if lista:
            print(f"{titulo}: {len(lista)}")
            for d in lista:
                print(f"  - {d}")
    falha = grava_se_conferido(saida, linhas, FONTE)
    if len(falhas) > falhas_esperadas:
        print(f"{len(falhas)} planilha(s) ilegivel(is), acima das {falhas_esperadas} esperada(s).", file=sys.stderr)
        return 1
    return falha


def main() -> int:
    p = argparse.ArgumentParser(description="Coleta as planilhas de cotacao da CEASA-MS e gera CSV.")
    p.add_argument("--start-year", type=int, default=ANO_INICIAL_PADRAO, help=f"padrao: {ANO_INICIAL_PADRAO}")
    p.add_argument("--end-year", type=int, default=date.today().year, help="padrao: ano corrente")
    p.add_argument("-b", "--bruto", type=Path, default=BRUTO_PADRAO, help="pasta das planilhas")
    p.add_argument("-s", "--saida", type=Path, default=SAIDA_PADRAO, help="CSV de saida")
    p.add_argument("--delay", type=float, default=0.5, help="segundos entre requisicoes (padrao: 0.5)")
    p.add_argument("--so-extrair", action="store_true", help="nao acessa o site, so refaz o CSV")
    p.add_argument("--falhas-esperadas", type=int, default=0, metavar="N",
                   help="planilhas ilegiveis toleradas antes de sair com erro (padrao: 0)")
    argumentos_incrementais(p)
    args = p.parse_args()

    erros = 0
    try:
        if not args.so_extrair:
            erros = coletar(Cliente(args.delay), args.bruto, args.start_year, args.end_year, args.meses_recentes)
    except requests.RequestException as e:
        # site fora do ar: o CSV ainda e gerado com o que ja foi baixado
        print(f"! coleta interrompida por falha de rede: {e}", file=sys.stderr)
        erros = 1
    falha = extrair(args.bruto, args.saida, args.falhas_esperadas, args.carregados, args.meses_recentes)
    return 2 if erros else falha


if __name__ == "__main__":
    sys.exit(main())
